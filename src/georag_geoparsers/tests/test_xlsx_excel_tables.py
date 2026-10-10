"""An .xlsx sheet is read whole - including rows below an Excel Table (finding 7).

``pl.read_excel(..., engine="openpyxl")`` reads ONLY the first Excel Table's
``ref`` when the sheet defines one (verified on polars 1.44.2: a table over
A1:D4 with six data rows came back with three). A re-export that appended
rows past the Table's last row, or a second block beside it, therefore lost
those rows with no warning - while ``enumerate_sheets`` (which counts every
row) still said the sheet had them. ``xlsx_parser`` now reads the sheet's rows
through openpyxl directly.

Also pinned here: the same cell-by-type text conversion the .xls path has
(finding 2) - a whole number is "1001" not "1001.0", Excel errors are blank
and reported, dates are ISO, and a repeated header is kept.
"""

from __future__ import annotations

import datetime
import re
import zipfile
from pathlib import Path

import pytest

openpyxl = pytest.importorskip("openpyxl")

from openpyxl.worksheet.table import Table  # noqa: E402

from georag_geoparsers.xlsx_parser import (  # noqa: E402
    enumerate_sheets,
    parse_xlsx_sheet,
    read_sheet_rows,
)

_HEADER = ["HoleID", "Easting", "Northing", "Total_Depth"]
_ROWS = [
    ["DH-1", 500000, 6000000, 100],
    ["DH-2", 500010, 6000010, 110],
    ["DH-3", 500020, 6000020, 120],
    ["DH-4", 500030, 6000030, 130],
    ["DH-5", 500040, 6000040, 140],
    ["DH-6", 500050, 6000050, 150],
]


def _workbook(path: Path, *, table_ref: str | None, rows=_ROWS, title: str = "Collars") -> Path:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = title
    ws.append(_HEADER)
    for row in rows:
        ws.append(row)
    if table_ref:
        ws.add_table(Table(displayName="CollarTable", ref=table_ref))
    wb.save(path)
    return path


class TestRowsPastAnExcelTable:
    def test_every_row_is_read_when_the_table_is_shorter_than_the_data(
        self, tmp_path: Path,
    ) -> None:
        # The Table covers the header and the first three data rows only.
        path = _workbook(tmp_path / "t.xlsx", table_ref="A1:D4")

        result = parse_xlsx_sheet(str(path), "Collars", "collar")

        assert [r["hole_id"] for r in result.records] == [f"DH-{i}" for i in range(1, 7)]
        assert result.total_rows == 6 and result.valid_rows == 6

    def test_a_table_that_covers_everything_still_reads_everything(
        self, tmp_path: Path,
    ) -> None:
        path = _workbook(tmp_path / "t.xlsx", table_ref="A1:D7")

        assert parse_xlsx_sheet(str(path), "Collars", "collar").valid_rows == 6

    def test_the_row_count_enumerate_reports_is_the_row_count_parsed(
        self, tmp_path: Path,
    ) -> None:
        path = _workbook(tmp_path / "t.xlsx", table_ref="A1:D4")

        meta = enumerate_sheets(str(path))[0]
        result = parse_xlsx_sheet(str(path), meta.name, "collar")

        assert meta.row_count == 6
        assert result.total_rows == meta.row_count

    def test_read_sheet_rows_reads_past_the_table_too(self, tmp_path: Path) -> None:
        path = _workbook(tmp_path / "t.xlsx", table_ref="A1:D4")

        rows = read_sheet_rows(str(path), "Collars")

        assert [r["HoleID"] for r in rows] == [f"DH-{i}" for i in range(1, 7)]
        assert rows[5]["Total_Depth"] == 150

    def test_a_wrong_dimension_record_does_not_cut_rows_off(self, tmp_path: Path) -> None:
        """Some exporters write a stale ``<dimension>``; openpyxl's read-only
        mode trusts it unless the dimensions are reset."""
        good = _workbook(tmp_path / "good.xlsx", table_ref=None)
        stale = tmp_path / "stale.xlsx"
        with zipfile.ZipFile(good) as src, zipfile.ZipFile(stale, "w") as dst:
            for item in src.infolist():
                data = src.read(item.filename)
                if item.filename == "xl/worksheets/sheet1.xml":
                    # `<dimension ref="A1:D7"/>` when openpyxl writes through
                    # lxml, `<dimension ref="A1:D7" />` through the stdlib
                    # writer it falls back to without it.
                    data, replaced = re.subn(
                        rb'<dimension ref="A1:D7"\s*/>', b'<dimension ref="A1:B2"/>', data
                    )
                    assert replaced == 1, data[:300]
                dst.writestr(item, data)

        result = parse_xlsx_sheet(str(stale), "Collars", "collar")

        assert result.valid_rows == 6
        assert result.records[5]["total_depth"] == 150.0


class TestCellsAreReadByType:
    def test_whole_numbers_and_dates_become_the_text_the_parser_reads(
        self, tmp_path: Path,
    ) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Collars"
        ws.append(["HoleID", "Easting", "Northing", "Elevation", "DrillDate"])
        ws.append([1001.0, 500000.5, 6000000.25, 300, datetime.datetime(2023, 4, 5)])
        ws.append([10010, 500100, 6000100, 310.0, datetime.datetime(2022, 11, 30, 8, 15)])
        path = tmp_path / "c.xlsx"
        wb.save(path)

        result = parse_xlsx_sheet(str(path), "Collars", "collar")

        by_id = {r["hole_id"]: r for r in result.records}
        # A float-typed 1001.0 is the hole "1001", not "1001.0" (whose canonical
        # form would collide with hole 10010).
        assert sorted(by_id) == ["1001", "10010"]
        assert by_id["1001"]["hole_id_canonical"] != by_id["10010"]["hole_id_canonical"]
        assert by_id["1001"]["elevation"] == 300.0
        assert by_id["10010"]["elevation"] == 310.0
        assert by_id["1001"]["drill_date"] == datetime.date(2023, 4, 5)
        assert by_id["10010"]["drill_date"] == datetime.date(2022, 11, 30)

    def test_excel_errors_are_blank_and_reported_not_kept_as_text(
        self, tmp_path: Path,
    ) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Collars"
        ws.append(["HoleID", "Easting", "Northing", "Elevation"])
        ws.append(["DH-1", 500000, 6000000, "#N/A"])
        ws.append(["DH-2", 500100, 6000100, 310])
        path = tmp_path / "c.xlsx"
        wb.save(path)

        result = parse_xlsx_sheet(str(path), "Collars", "collar")

        by_id = {r["hole_id"]: r for r in result.records}
        assert by_id["DH-1"]["elevation"] is None
        assert by_id["DH-2"]["elevation"] == 310.0
        note = next(w for w in result.warnings if w["code"] == "excel_error_cells")
        assert note["context"]["cells"] == [{"cell": "D2", "error": "#N/A"}]
        assert note["detail"]

    def test_a_repeated_header_is_kept_under_a_distinct_name(self, tmp_path: Path) -> None:
        wb = openpyxl.Workbook()
        ws = wb.active
        ws.title = "Collars"
        ws.append(["HoleID", "Easting", "Northing", "Notes", "Notes"])
        ws.append(["DH-1", 500000, 6000000, "first", "second"])
        path = tmp_path / "c.xlsx"
        wb.save(path)

        rows = read_sheet_rows(str(path), "Collars")
        result = parse_xlsx_sheet(str(path), "Collars", "collar")

        assert rows == [{
            "HoleID": "DH-1", "Easting": 500000, "Northing": 6000000,
            "Notes": "first", "Notes_4": "second",
        }]
        renamed = next(w for w in result.warnings if w["code"] == "duplicate_header_renamed")
        assert renamed["severity"] == "info"
        assert renamed["context"]["renamed"] == [("Notes", "Notes_4")]
